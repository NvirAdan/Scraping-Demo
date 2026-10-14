"""Export cleaned rows to CSV and Excel."""

import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Font


def export_csv(rows, path):
    """Write rows to CSV (utf-8-sig so Excel opens accents correctly)."""
    if not rows:
        print(f"[-] Nothing to export to {path}")
        return
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=rows[0].keys())
        writer.writeheader()
        writer.writerows(rows)
    print(f"[+] CSV saved: {path} ({len(rows)} rows)")


def export_xlsx(rows, path):
    """Write rows to a formatted Excel file."""
    if not rows:
        print(f"[-] Nothing to export to {path}")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"

    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row[h] for h in headers])

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        longest = max(len(str(r[header])) for r in rows)
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = (
            min(max(len(header), longest) + 2, 60)
        )

    wb.save(path)
    print(f"[+] Excel saved: {path} ({len(rows)} rows)")


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)
